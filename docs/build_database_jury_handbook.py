from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SCHEMA = ROOT / "database" / "schema.sql"
OUTPUT = ROOT / "docs" / "Database_Design_and_Jury_Preparation_Handbook.docx"

NAVY = "173F5F"
BLUE = "DCEAF4"
PALE = "F3F7FA"
GRID = "D9D9D9"
TEXT = RGBColor(35, 43, 49)


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), GRID)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])


def add_table(doc, headers, rows, widths=None, font_size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        set_cell_fill(cell, NAVY)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(font_size)
        cell.paragraphs[0].paragraph_format.keep_with_next = True
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(values):
            cells[col_index].text = str(value)
            cells[col_index].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cells[col_index])
            if row_index % 2:
                set_cell_fill(cells[col_index], PALE)
            for paragraph in cells[col_index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(font_size)
                    run.font.color.rgb = TEXT
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbered(doc, items):
    for index, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.first_line_indent = Inches(-0.22)
        p.add_run(f"{index}. ").bold = True
        p.add_run(item)


def add_code(doc, code):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.right_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(code)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    run.font.size = Pt(9)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    p_pr.append(shd)


def extract_foreign_keys():
    text = SCHEMA.read_text(encoding="utf-8")
    current = None
    rows = []
    table_pattern = re.compile(r"^CREATE TABLE\s+`?([A-Za-z0-9_]+)`?", re.I)
    fk_pattern = re.compile(
        r"CONSTRAINT\s+([A-Za-z0-9_]+)\s+FOREIGN KEY\s*\(([^)]+)\)\s+"
        r"REFERENCES\s+`?([A-Za-z0-9_]+)`?\s*\(([^)]+)\)(.*)", re.I
    )
    for line in text.splitlines():
        match = table_pattern.search(line)
        if match:
            current = match.group(1)
        match = fk_pattern.search(line)
        if match:
            tail = match.group(5).upper()
            delete_rule = "CASCADE" if "ON DELETE CASCADE" in tail else "SET NULL" if "ON DELETE SET NULL" in tail else "RESTRICT"
            rows.append((current, match.group(2).replace("`", "").strip(), match.group(3), match.group(4).replace("`", "").strip(), delete_rule))
    return rows


TABLES = [
    ("Security and governance", "roles", "Role definitions and permission JSON", "id", "Parent of users"),
    ("Security and governance", "users", "Authenticated application accounts", "id", "References roles"),
    ("Academic master", "schools", "University schools", "id", "Parent of departments, programmes and faculty"),
    ("Academic master", "departments", "Departments within a school", "id", "References schools"),
    ("Academic master", "programmes", "Branches and academic programmes", "id", "References schools and departments"),
    ("Academic master", "batches", "Admission cohorts such as UCS 2023-2027", "id", "References programmes"),
    ("Academic master", "courses", "Reusable subject master", "id", "Referenced by batch curriculum and examinations"),
    ("Academic master", "programme_courses", "Batch-wise curriculum by semester and section", "id", "References programme, batch and course"),
    ("People", "students", "Student identity and academic placement", "id", "References programme and batch"),
    ("People", "faculty", "Faculty master and organisational placement", "id", "References school and department"),
    ("Infrastructure", "rooms", "Examination rooms and capacities", "id", "Parent of room seats"),
    ("Infrastructure", "room_seats", "Visual seat coordinates and enabled state", "id", "References rooms; cascades on room deletion"),
    ("Exam planning", "exam_cycles", "An examination window for one academic session", "id", "Created by a user"),
    ("Exam planning", "exam_cycle_batches", "Batches selected for an exam cycle", "cycle_id + batch_id", "Junction between cycles and batches"),
    ("Exam planning", "exam_shifts", "Morning, evening or custom sessions", "id", "References exam cycles"),
    ("Exam planning", "exam_calendar_dates", "Enabled dates, Sundays and holidays", "id", "References exam cycles"),
    ("Automatic scheduling", "scheduling_rules", "School-specific gap and priority configuration", "id", "References cycle, school and user"),
    ("Automatic scheduling", "scheduling_runs", "Validation, generation and publication attempts", "id", "References cycle, school and users"),
    ("Automatic scheduling", "scheduling_run_items", "Selected curriculum paper and assigned slot", "id", "References run, curriculum, shift and exam"),
    ("Automatic scheduling", "scheduling_conflicts", "Blocking errors and warnings detected by validation", "id", "References scheduling run"),
    ("Date sheet", "examinations", "Scheduled written paper on a date and shift", "id", "References cycle, shift, course and run"),
    ("Date sheet", "examination_cohorts", "Programme, batch, semester and section rows for a paper", "id", "References examination, programme and batch"),
    ("Date sheet", "exam_eligibility", "Student-level eligibility for a paper", "id", "References examination and student"),
    ("Seating", "seating_allocations", "Versioned seating generation for a session", "id", "References cycle, shift and generator"),
    ("Seating", "seating_assignments", "Student to room and seat assignment", "id", "References allocation, exam, room, seat and student"),
    ("Seating", "seating_unallocated", "Students that could not receive a seat", "id", "References allocation, exam and student"),
    ("Invigilation", "faculty_availability", "Faculty availability by date and shift", "id", "References faculty and shift"),
    ("Invigilation", "invigilation_allocations", "Room duty assigned to faculty", "id", "References cycle, shift, room, faculty and user"),
    ("Invigilation", "replacement_requests", "Replacement workflow for an invigilation duty", "id", "References duty, users and replacement faculty"),
    ("Attendance", "attendance", "Student attendance for an examination", "id", "References seating, examination, room, student and users"),
    ("Imports", "import_batches", "Uploaded bulk file and processing summary", "id", "References uploader"),
    ("Imports", "import_rows", "Normalized staged rows before commit", "id", "References import batch"),
    ("Imports", "import_errors", "Row and column validation errors", "id", "References import batch"),
    ("Governance", "audit_logs", "Immutable record of important actions", "id", "References user; entity reference is polymorphic"),
    ("Configuration", "system_settings", "Application-wide typed settings", "setting_key", "References updating user"),
]


JURY_QA = [
    ("What problem does the database solve?", "It provides one consistent source of truth for academic masters, batch-wise curriculum, exam cycles, automatic date-sheet generation, seating, invigilation, attendance, imports and audit history."),
    ("Why is the database relational?", "The data contains stable relationships such as school to programme, programme to batch and examination to student. Foreign keys, unique constraints and transactions protect these relationships better than duplicated files or isolated spreadsheets."),
    ("Why are courses and programme_courses separate?", "courses stores the reusable subject identity. programme_courses stores where that subject applies: programme, batch, semester, section, credits, duration and priority. This avoids duplicating course names while supporting different curricula."),
    ("How is batch-wise curriculum implemented?", "programme_courses contains both programme_id and batch_id. A unique constraint prevents the same course from being mapped twice to the same programme, batch, semester and section."),
    ("Why does batch belong to programme?", "A label such as UCS 2023-2027 is meaningful only in the context of UCS. programme_id guarantees that the cohort cannot be confused with a similarly dated cohort in another programme."),
    ("How is programme duration stored?", "programmes.duration_semesters stores the total duration. Batch end year can be derived from the admission year and duration, while the explicit batch dates preserve the official cohort label."),
    ("How does the system determine the current semester?", "It compares the batch admission period with the examination academic session and maps elapsed academic terms to a semester, bounded by programme duration. Completed batches are excluded from regular scheduling."),
    ("How are laboratory subjects excluded?", "The application applies the written-paper policy to course category and naming/type information before creating scheduling items. Only eligible written curriculum records proceed to scheduling."),
    ("Why is examination_cohorts required?", "One scheduled course paper may serve one or more programme-batch-semester-section cohorts. This table preserves each cohort separately instead of combining branches into an ambiguous text field."),
    ("Can different branches have papers at the same time?", "Yes. Different cohorts may share a date and shift when no defined conflict exists. The schedule is cohort-aware, so unrelated branches do not unnecessarily block each other."),
    ("How are subject priorities used?", "programme_courses.subject_priority provides an order from high to low importance. The scheduling rule can place higher-priority written subjects earlier while still respecting dates, gaps and shifts."),
    ("How are holidays represented?", "exam_calendar_dates stores each date with enabled and holiday flags. The generator schedules only enabled examination dates and can account for Sundays and explicitly configured holidays."),
    ("What is a scheduling run?", "It is a traceable attempt to validate or generate a schedule for a school and cycle. It stores configuration, status, creator, approver and publisher, while run items and conflicts preserve the detailed result."),
    ("Why are scheduling conflicts stored?", "Persisted conflicts make validation explainable. The UI can show exactly why a run is blocked or warned, and the audit trail remains available for review."),
    ("How is a generated exam linked back to automation?", "examinations.generated_by_run_id references scheduling_runs. If the run is removed, the reference becomes NULL so a valid published examination is not accidentally deleted."),
    ("What prevents duplicate examinations?", "A unique constraint covers cycle, course, exam date, shift and category. Cohort uniqueness is enforced separately in examination_cohorts."),
    ("How does seating remain versioned?", "seating_allocations stores cycle, date, shift and version number. Assignments belong to a specific version, allowing regeneration without overwriting earlier results."),
    ("What prevents two students from receiving the same seat?", "seating_assignments has a unique constraint on allocation, room and seat. It also uniquely constrains allocation and student, preventing duplicate placement."),
    ("How are disabled seats handled?", "room_seats stores an enabled flag. The allocation service selects only enabled seats, so blocked or damaged seats remain visible in the layout but are never assigned."),
    ("How is faculty double-booking prevented?", "Unique duty and availability constraints use faculty, exam date and shift. The allocation logic checks these records before saving an invigilation duty."),
    ("Why is attendance linked to examination and student?", "That pair is the real business identity of an attendance result. A unique constraint ensures one final attendance record per student per examination."),
    ("How do bulk imports work safely?", "A file creates an import_batches record. Every row is staged in import_rows, validation problems go to import_errors, and only valid rows are committed within a transaction."),
    ("Why is source_row_number used?", "It preserves the original CSV row for understandable error reporting and prevents duplicate staged rows inside one import batch."),
    ("What is referential integrity?", "It is the guarantee that a foreign-key value points to an existing parent row. For example, a student cannot reference a programme or batch that does not exist."),
    ("What is the difference between primary, foreign and unique keys?", "A primary key identifies a row. A foreign key links it to another table. A unique key prevents duplicate business combinations such as one seat being assigned twice in an allocation."),
    ("What does ON DELETE CASCADE mean?", "When a parent is deleted, dependent detail records are deleted automatically. It is used where child data has no independent meaning, such as room seats belonging to a room."),
    ("Why is deletion restricted in many tables?", "Historical exam, attendance and allocation data must not become orphaned. Restriction forces the application to review or remove dependent records in a controlled sequence."),
    ("What is normalization in this schema?", "The design largely follows third normal form: each table represents one subject, repeating groups are separated, and non-key attributes depend on the table key rather than unrelated data."),
    ("Give an example of avoiding redundancy.", "A course name is stored once in courses. programme_courses references it and adds curriculum context instead of copying the name for every batch and semester."),
    ("Where is denormalization used?", "Operational snapshot or display fields may be retained where historical output must remain stable. Such fields are intentional and do not replace the authoritative master relationships."),
    ("Why are indexes necessary?", "Indexes reduce lookup cost for frequent filters such as programme, batch, semester, section, date and shift. They also support unique business rules."),
    ("What transaction boundaries are important?", "Import commit, date-sheet generation, seating generation and multi-table deletion should be transactional. Either every dependent write succeeds or the entire operation rolls back."),
    ("How is role-based access represented?", "roles stores permission definitions and users references one role. The application checks permissions before controllers allow operations such as publishing, importing or deleting."),
    ("How does audit logging differ from application logs?", "audit_logs records business actions, actor and affected entity. Server logs diagnose technical behaviour; audit records explain who changed business data."),
    ("Why is audit entity_id not a foreign key?", "It is a polymorphic reference that may identify different table types. Keeping it independent also preserves the audit record after the original entity is removed."),
    ("How is SQL injection prevented?", "Application queries use prepared statements with bound parameters. User input is never concatenated directly into executable SQL."),
    ("How are passwords stored?", "Only secure password hashes are stored in users.password_hash. Plain-text passwords must never be written to the database or seed file."),
    ("What is the purpose of schema.sql?", "It is the canonical fresh-install definition containing every table, constraint and index required by the application."),
    ("What is the purpose of seed.sql?", "It inserts repeatable baseline data such as roles, schools, programmes, initial batches, course curriculum and system settings after schema creation."),
    ("Is seed.sql a backup?", "No. It is deterministic starter data. User-created cycles, students, seating and operational records require a database backup or export."),
    ("In what order should a fresh database be installed?", "Create/select the target database, import schema.sql first, then import seed.sql, configure the environment credentials, and verify the health endpoint and login."),
    ("How is the hosted database different from local XAMPP?", "The application model is the same, but hosted credentials use the provider's database name, username and hostname. Shared hosts may restrict CREATE DATABASE and temporary-table privileges."),
    ("How would you back up the database?", "Use phpMyAdmin export or mysqldump with structure and data, verify the file, retain dated copies, and test restoration into a separate database."),
    ("How do you prove the schema and seed are compatible?", "The project creates a temporary fresh database from schema.sql, runs seed.sql twice, and verifies expected counts. The repeatability test currently passes."),
    ("What are the main future database improvements?", "Potential improvements include formal migrations, automated backup retention, read-only reporting replicas for scale, and more granular permission assignments if institutional roles expand."),
]


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].font.color.rgb = TEXT
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.08
    for name, size in (("Title", 28), ("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 11.5)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    header = section.header.paragraphs[0]
    header.text = "Exam Scheduling and Management System  |  Database Handbook"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(90, 102, 110)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Database Design and Jury Preparation  |  Page ")
    add_page_number(footer)
    for run in footer.runs:
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(90, 102, 110)

    # Cover
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(110)
    p.add_run("Database Design and Jury Preparation Handbook")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Exam Scheduling and Management System")
    r.bold = True
    r.font.size = Pt(16)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Gautam Buddha University Examination Operations")
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(70, 85, 94)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Prepared from the current application schema").italic = True
    doc.add_page_break()

    doc.add_heading("How to Use This Handbook", level=1)
    doc.add_paragraph(
        "This handbook explains the database in presentation-ready language. Read Sections 1 to 7 to understand the design, use Sections 8 to 11 during demonstration and troubleshooting, and revise the question bank before the jury. The current implementation contains 35 tables and 73 declared foreign keys. Its central design decision is to keep curriculum and scheduling separated by programme, admission batch, semester and section."
    )
    add_bullets(doc, [
        "For a five-minute explanation, study the executive summary, core relationship chain and scheduling data flow.",
        "For a technical jury, focus on keys, normalization, constraints, transactions, indexes and deletion behaviour.",
        "For a live demonstration, follow the recommended data-entry order and verification queries.",
        "For deployment, import schema.sql before seed.sql and never treat seed.sql as a production backup.",
    ])

    doc.add_heading("Contents", level=1)
    contents = [
        "Database Overview", "Relational Design Concepts", "Complete Table Dictionary",
        "Core Entity Relationships", "Automatic Scheduling Data Model", "Seating Invigilation and Attendance",
        "Integrity Normalization and Performance", "Database Installation and Maintenance",
        "Demonstration and Verification Queries", "Troubleshooting Guide", "Jury Question Bank",
        "Appendix A Complete Foreign Key Matrix", "Appendix B Presentation Script and Revision Checklist"
    ]
    add_numbered(doc, contents)
    doc.add_page_break()

    doc.add_heading("1 Database Overview", level=1)
    doc.add_paragraph(
        "The database supports the complete examination workflow rather than only storing a date sheet. Academic master data defines schools, departments, programmes, batches, subjects, students, faculty and rooms. Examination planning then selects batches, calendar dates and shifts. Automatic scheduling creates traceable runs and conflict records. Published examinations feed eligibility, seating, invigilation and attendance. Import and audit tables make administrative actions reviewable."
    )
    add_table(doc, ["Measure", "Current design"], [
        ("Database engine", "MySQL or MariaDB using InnoDB"),
        ("Tables", "35"),
        ("Declared foreign keys", "73"),
        ("Canonical installation files", "database/schema.sql then database/seed.sql"),
        ("Primary scheduling scope", "School, programme, batch, semester, section and selected written subjects"),
        ("Major operational outputs", "Date sheet, seating plan, invigilation duties, attendance and reports"),
    ], widths=[2.2, 4.7])

    doc.add_heading("System Data Flow", level=2)
    add_numbered(doc, [
        "Create schools and departments, then programmes with their duration.",
        "Create admission batches for each programme and enrol students into the correct batch.",
        "Create reusable courses and map them to a programme and batch through programme_courses.",
        "Create an exam cycle, select eligible batches, configure shifts, dates, holidays and rules.",
        "Validate and generate a scheduling run; review conflicts before creating examinations.",
        "Attach programme and batch cohorts to each examination and determine student eligibility.",
        "Generate versioned seating, assign invigilation, record attendance and retain the audit history.",
    ])

    doc.add_heading("2 Relational Design Concepts", level=1)
    concepts = [
        ("Entity", "A real system concept represented by a table, such as students or exam_cycles."),
        ("Attribute", "A property represented by a column, such as roll_number or exam_date."),
        ("Primary key", "A value that uniquely identifies a row. Most tables use id."),
        ("Foreign key", "A column that must reference an existing row in another table."),
        ("Unique key", "A rule preventing duplicate business data or duplicate combinations."),
        ("Junction table", "A table resolving a many-to-many relationship, such as exam_cycle_batches."),
        ("Index", "A database structure that accelerates filtering and joining."),
        ("Transaction", "A group of writes that either all commit or all roll back."),
    ]
    add_table(doc, ["Term", "Meaning in this project"], concepts, widths=[1.55, 5.35])

    doc.add_heading("Relationship Types Used", level=2)
    add_bullets(doc, [
        "One to many: one school offers many programmes.",
        "Many to many through a junction: an exam cycle selects many batches and a batch may appear in many cycles.",
        "Identifying dependency: room seats cannot exist without a room.",
        "Optional historical reference: a generated examination can remain after its scheduling run reference is set to NULL.",
    ])

    doc.add_heading("3 Complete Table Dictionary", level=1)
    doc.add_paragraph("Every table in the current schema is listed below. The purpose column is the simplest sentence to use when the jury asks what the table stores.")
    grouped = {}
    for domain, table, purpose, key, relation in TABLES:
        grouped.setdefault(domain, []).append((table, purpose, key, relation))
    for domain, rows in grouped.items():
        doc.add_heading(domain, level=2)
        add_table(doc, ["Table", "Purpose", "Primary key", "Main relationship"], rows, widths=[1.45, 2.65, 1.15, 1.75], font_size=8.3)

    doc.add_heading("4 Core Entity Relationships", level=1)
    doc.add_heading("Academic hierarchy", level=2)
    add_code(doc, "SCHOOL -> DEPARTMENT -> PROGRAMME -> BATCH -> STUDENT\n                    PROGRAMME + BATCH + COURSE -> PROGRAMME_COURSE")
    doc.add_paragraph(
        "The programme is the branch, while the batch is the admission cohort. A student belongs to one programme and one batch. A course is a reusable subject master. programme_courses supplies the missing context by stating that a particular course belongs to a particular programme and batch in a specific semester and section."
    )
    doc.add_heading("Why programme and batch must not be merged", level=2)
    doc.add_paragraph(
        "UCS describes B.Tech Computer Science and Engineering, but UCS 2023-2027 identifies one admission cohort. Keeping these in separate tables supports multiple active batches, completed batches, lateral-entry durations and curriculum changes between admission years."
    )
    doc.add_heading("Important business uniqueness rules", level=2)
    add_table(doc, ["Rule", "Purpose"], [
        ("department school and code", "Department code is unique within its school."),
        ("batch programme and label", "A programme cannot contain duplicate batch labels."),
        ("programme batch course semester section", "A batch curriculum mapping cannot be duplicated."),
        ("room and seat label", "A room cannot contain duplicate seat numbers."),
        ("cycle and shift sequence", "Shift order remains unique inside a cycle."),
        ("cycle and calendar date", "A date is configured once per cycle."),
    ], widths=[2.9, 4.0])

    doc.add_heading("5 Automatic Scheduling Data Model", level=1)
    doc.add_heading("Scheduling scope", level=2)
    doc.add_paragraph(
        "The generator works school-wise but keeps every selected programme, batch, semester and section as a distinct cohort. Courses do not become one combined branch row. Subject filters affect only the selected curriculum mappings and retain the user's deselection state."
    )
    add_code(doc, "EXAM_CYCLE -> SELECTED_BATCHES + CALENDAR_DATES + SHIFTS\nPROGRAMME_COURSES -> SCHEDULING_RUN -> RUN_ITEMS -> EXAMINATIONS -> EXAMINATION_COHORTS")
    doc.add_heading("Validation sequence", level=2)
    add_numbered(doc, [
        "Confirm selected programmes belong to the chosen school.",
        "Resolve the active/current semester for each selected programme batch.",
        "Load active batch-wise curriculum and exclude labs or practical subjects.",
        "Apply explicit subject selections and priorities.",
        "Check enabled dates, holidays, shifts and requested gaps.",
        "Evaluate capacity for each cohort rather than imposing a fixed example limit.",
        "Store blockers and warnings in scheduling_conflicts.",
        "Generate examinations and examination_cohorts only after a valid run is saved.",
    ])
    doc.add_heading("Gap and holiday logic", level=2)
    doc.add_paragraph(
        "A preferred gap is applied when calendar capacity allows it. If a cohort has as many papers as available working days, the generator must use consecutive enabled days rather than fail. Sundays and configured holidays are never selected as examination dates. Different branches may use the same date and shift when they do not represent the same cohort conflict."
    )

    doc.add_heading("6 Seating Invigilation and Attendance", level=1)
    doc.add_heading("Seating", level=2)
    doc.add_paragraph(
        "Rooms contain individually addressable seats. Disabled seats remain part of the visual room layout but are excluded from allocation. Each seating generation creates a versioned seating_allocations row; assignments and failures belong to that version."
    )
    doc.add_heading("Invigilation", level=2)
    doc.add_paragraph(
        "Faculty availability is recorded by date and shift. Invigilation duties connect a cycle, shift, room and faculty member. Replacement requests preserve the original duty and the review trail rather than directly overwriting the assignment."
    )
    doc.add_heading("Attendance", level=2)
    doc.add_paragraph(
        "Attendance connects the examination, student, room and seating allocation. The examination and student combination is unique, ensuring that one candidate cannot receive multiple final attendance states for one paper."
    )

    doc.add_heading("7 Integrity Normalization and Performance", level=1)
    doc.add_heading("Normalization", level=2)
    add_table(doc, ["Normal form idea", "How the design applies it"], [
        ("First normal form", "Columns contain atomic values and repeating records use separate rows."),
        ("Second normal form", "Junction-table attributes describe the full relationship key."),
        ("Third normal form", "Schools, programmes, courses, batches and people are stored once and referenced by keys."),
    ], widths=[1.9, 5.0])
    doc.add_heading("Referential actions", level=2)
    add_bullets(doc, [
        "CASCADE is used for owned detail data, including room seats, selected cycle batches, calendar dates, run details and import staging rows.",
        "SET NULL is used when a historical examination or run item should survive without its former generator reference.",
        "RESTRICT is the safe default for operational records that require deliberate dependency review before deletion.",
    ])
    doc.add_heading("Indexes and query performance", level=2)
    doc.add_paragraph(
        "The schema indexes common academic filters, date and shift lookups, subject priority, audit entities and import rows. Unique constraints also create indexes. For large deployments, query plans should be reviewed with EXPLAIN before adding new indexes because excessive indexing slows imports and updates."
    )
    doc.add_heading("Transaction requirements", level=2)
    add_bullets(doc, [
        "Commit bulk imports only when all intended valid rows can be persisted.",
        "Create a scheduling run, its items and resulting examinations atomically.",
        "Generate seating assignments and the unallocated list in one transaction.",
        "Delete complex exam-cycle history through a service that removes dependants in a verified order.",
    ])

    doc.add_heading("8 Database Installation and Maintenance", level=1)
    doc.add_heading("Fresh installation", level=2)
    add_numbered(doc, [
        "Create or select an empty MySQL database.",
        "Import database/schema.sql.",
        "Import database/seed.sql.",
        "Set DB_HOST, DB_PORT, DB_DATABASE, DB_USERNAME and DB_PASSWORD in the environment file.",
        "Open the health endpoint and confirm database connected.",
        "Sign in and verify roles, programmes, batches and curriculum before creating a cycle.",
    ])
    doc.add_paragraph("For an existing database, do not blindly re-import the full schema because CREATE TABLE statements are intended for a fresh installation. Back up first and apply a reviewed migration when structural changes are required.")
    doc.add_heading("Schema and seed responsibilities", level=2)
    add_table(doc, ["File", "Contains", "Does not replace"], [
        ("schema.sql", "Tables, columns, keys, indexes and constraints", "A migration plan for an already populated database"),
        ("seed.sql", "Repeatable baseline roles, masters, curriculum and settings", "A backup of operational data created in the UI"),
    ], widths=[1.3, 3.0, 2.6])
    doc.add_heading("Backup practice", level=2)
    add_bullets(doc, [
        "Export both structure and data before deployments or major imports.",
        "Use dated backup names and retain more than one generation.",
        "Do not store production credentials or raw backups in the public web directory or Git repository.",
        "Test restoration into a separate database; an untested backup is not a verified recovery plan.",
    ])

    doc.add_heading("9 Demonstration and Verification Queries", level=1)
    doc.add_paragraph("These read-only queries are suitable for a jury demonstration. Replace numeric IDs only when necessary.")
    queries = [
        ("Confirm selected database", "SELECT DATABASE();"),
        ("Count application tables", "SELECT COUNT(*) AS total_tables FROM information_schema.tables WHERE table_schema = DATABASE() AND table_type = 'BASE TABLE';"),
        ("Show programme batches", "SELECT p.code, p.name, b.label, b.start_year, b.end_year, b.status FROM programmes p JOIN batches b ON b.programme_id = p.id ORDER BY p.code, b.start_year;"),
        ("Show batch curriculum", "SELECT p.code, b.label, pc.semester, c.code AS course_code, c.name, pc.subject_priority FROM programme_courses pc JOIN programmes p ON p.id = pc.programme_id JOIN batches b ON b.id = pc.batch_id JOIN courses c ON c.id = pc.course_id ORDER BY p.code, b.start_year, pc.semester, pc.subject_priority;"),
        ("Show scheduled cohorts", "SELECT e.exam_date, s.name AS shift_name, c.code AS course_code, p.code AS programme_code, b.label, ec.semester, ec.section FROM examinations e JOIN exam_shifts s ON s.id = e.shift_id JOIN courses c ON c.id = e.course_id JOIN examination_cohorts ec ON ec.examination_id = e.id JOIN programmes p ON p.id = ec.programme_id JOIN batches b ON b.id = ec.batch_id ORDER BY e.exam_date, s.sequence_no, p.code, b.start_year;"),
        ("Check seating duplicates", "SELECT allocation_id, room_id, seat_id, COUNT(*) AS uses FROM seating_assignments GROUP BY allocation_id, room_id, seat_id HAVING COUNT(*) > 1;"),
        ("Check latest imports", "SELECT id, import_type, file_name, total_rows, valid_rows, invalid_rows, status FROM import_batches ORDER BY id DESC LIMIT 10;"),
        ("Check latest audit actions", "SELECT created_at, user_id, action, entity_type, entity_id FROM audit_logs ORDER BY id DESC LIMIT 20;"),
    ]
    for title, query in queries:
        doc.add_heading(title, level=2)
        add_code(doc, query)

    doc.add_heading("10 Troubleshooting Guide", level=1)
    add_table(doc, ["Symptom", "Likely cause", "Safe check or correction"], [
        ("Unknown column source_row_number", "Existing import_rows table predates the current schema", "Inspect SHOW COLUMNS first; apply a targeted migration rather than repeating an invalid rename."),
        ("Foreign-key error while deleting", "Dependent operational records still reference the parent", "Identify child rows from the FK matrix and use the application deletion service or a reviewed transaction."),
        ("Only information_schema views appear", "phpMyAdmin is displaying the system database", "Select the application database and run SELECT DATABASE()."),
        ("Access denied for CREATE TEMPORARY TABLE", "Shared host restricts temporary-table privileges", "Use the hosting-compatible seed package or import a seed version that avoids restricted statements."),
        ("Login redirects to login", "Credentials, password hash, session storage or cookie configuration is invalid", "Confirm database connection, active user, password hash and production session/cookie settings."),
        ("No curriculum papers found", "The selected programme, batch and semester has no active programme_courses", "Verify the exact batch label and curriculum mapping with the batch curriculum query."),
        ("Schedule capacity warning", "Requested papers and gaps exceed enabled slots for a cohort", "Review enabled dates, shifts and preferred gap; do not treat unrelated branches as the same cohort."),
        ("Uploaded file cannot be stored", "Upload directory is missing or not writable", "Verify configured storage path, PHP upload limits and server permissions."),
    ], widths=[1.75, 2.35, 2.9], font_size=8.4)

    doc.add_heading("11 Jury Question Bank", level=1)
    doc.add_paragraph("Answer the first sentence directly, then add the second sentence only if the jury asks for implementation detail.")
    for index, (question, answer) in enumerate(JURY_QA, 1):
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        run = p.add_run(f"Q{index}. {question}")
        run.bold = True
        p = doc.add_paragraph(answer)
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.space_after = Pt(8)

    doc.add_heading("Appendix A Complete Foreign Key Matrix", level=1)
    fk_rows = extract_foreign_keys()
    doc.add_paragraph(f"The following matrix is generated from the current schema.sql and contains all {len(fk_rows)} declared foreign keys.")
    add_table(doc, ["Child table", "Foreign-key column", "Parent table", "Parent key", "On delete"], fk_rows, widths=[1.55, 1.75, 1.55, 0.85, 0.85], font_size=7.8)

    doc.add_page_break()
    doc.add_heading("Appendix B Presentation Script and Revision Checklist", level=1)
    doc.add_heading("Two minute database explanation", level=2)
    doc.add_paragraph(
        "Our application uses a normalized MySQL relational database with 35 tables and 73 foreign keys. The academic hierarchy starts with school, department and programme. Every programme has admission batches, and students belong to the correct programme and batch. Subjects are stored once in courses, while programme_courses creates a batch-wise curriculum for a semester and section. An exam cycle selects batches, dates and shifts. Automatic scheduling validates written subjects, holidays, gaps and priorities, then stores a scheduling run, its items and any conflicts. Valid items become examinations, and examination_cohorts keeps every branch and batch separate. Those examinations then drive eligibility, seating, invigilation and attendance. Foreign keys, unique constraints, transactions and audit logs protect data consistency and traceability."
    )
    doc.add_heading("Before the jury", level=2)
    checklist = [
        "I can explain the difference between programme, batch, course and programme_course.",
        "I can trace one student from programme and batch to examination, seat and attendance.",
        "I can explain PK, FK, UK, indexes, normalization and transactions with project examples.",
        "I can explain why examinations and examination_cohorts are separate.",
        "I can explain CASCADE, SET NULL and RESTRICT deletion behaviour.",
        "I can run the demonstration queries without modifying data.",
        "I know that schema.sql is for structure and seed.sql is baseline data, not a backup.",
        "I have a verified database backup before the demonstration.",
    ]
    for item in checklist:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run("[ ] " + item)

    # Prevent isolated headings where possible.
    for paragraph in doc.paragraphs:
        if paragraph.style.name.startswith("Heading"):
            keep_with_next(paragraph)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
